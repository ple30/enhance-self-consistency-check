from google import genai
from google.genai import types
import pandas as pd
from pathlib import Path
import re
from openpyxl import load_workbook

client = genai.Client(api_key="")
perturbation_instruction = f"""You are an expert at evaluating perturbation strength in sentences. Use the following rating guide to assess perturbed sentences:
For the following descriptions, reasoning path refers to the reasoning in the provided CoT.
Interfere with the reasoning path means that it changes, invalidates, or redirects the logical steps required to derive the answer. This can be done by introducing new premises that might contradict with the current reasoning, removing necessary premises, or altering logical dependencies.
When assessing the text, consider whether any unchanged premises or clauses still support the original answer. Give greater weight to unchanged conclusion clauses that maintain support when determining the final score.

Use the reasoning path provided to assess the edited text and choose one of the followings:
- No/Minimal change: Meaning is preserved (e.g., paraphrasing, synonyms, rewording) OR the change(s) is/are not enough to interfere with the reasoning path
- Small change: Minor semantic shift that might interfere with the reasoning path but not enough to remove or reverse support for an answer option
- Moderately strong change: The meaning changes enough to notably weaken, remove, or reverse support for an answer option, without clearly shifting support to another option (including cases where it shifts support to an answer that is not in the answer options).
- Strong change: Clearly shift support from one answer option to another that is in the option list. In the perturbed text, there should be only few or little support for original answer. If there is a strong support for another answer but still a very strong support for the original answer, it should be a moderately strong change.

RATING SCALE:
0: No/minimal change
1: small change
2: moderately strong change
3: strong change

Note: When the score seems to be in between two ratings, you should give a lower one
For the perturbed text provided, give:
- Rating (0-3)
- Brief justification"""

with open(Path(__file__).parent / 'examples_for_llm' / 'sqa_input_example_1.txt', "r") as f:
    examples = f.read()



def create_prompt_input(question, cot, edited_question, examples):
    return f"""Here are examples of how to apply the rating guide to a perturbed question:
{examples}

Now, rate the following:
Original Question: {question}

CoT: {cot}
Perturbed Question: {edited_question}
"""

def create_prompt_cot(question, cot, edited_cot, examples):
    return f"""Here are examples of how to apply the rating guide to a perturbed CoT:
{examples}

Now, rate the following:
Question: {question}

Original CoT: {cot}
Perturbed CoT: {edited_cot}
"""

def extract_rating(response_text):
    """Extract the number after 'Rating: ' from the response."""
    match = re.search(r'Rating:\s*([0-3])', response_text)
    if match:
        return int(match.group(1))
    return None  # Return None if no rating found


if __name__ == "__main__":
    input_folder = Path(__file__).parent / 'annotation_data'

    for input_file in input_folder.glob('sqa_input_annotation.xlsx'):
        print(f"\nProcessing file: {input_file.name}")

        sheets = pd.read_excel(input_file, sheet_name=None)
        wb = load_workbook(input_file)

        for sheet_name, df in sheets.items():
            print(f"  Sheet: {sheet_name}")
            ws = wb[sheet_name]

            # Filter rows where score is between 0 and 3
            df['score'] = pd.to_numeric(df['score'], errors='coerce')
            mask = df['score'].between(0, 3)
            filtered_df = df[mask].copy()

            if filtered_df.empty:
                print(f"  No valid rows, skipping.")
                continue

            # Find or create 'gemini' column
            header_row = [cell.value for cell in ws[1]]
            if 'gemini' in header_row:
                gemini_col_idx = header_row.index('gemini') + 1
            else:
                gemini_col_idx = len(header_row) + 1
                ws.cell(row=1, column=gemini_col_idx, value='gemini')

            for idx, row in filtered_df.iterrows():
                question = str(row['Question'])
                cot = str(row['Full original CoT'])
                edited_question = str(row['Edited Question'])
                #edited_cot = str(row['Full edited CoT'])

                prompt_for_input = create_prompt_input(question, cot, edited_question, examples)
                #prompt_for_input = create_prompt_cot(question, cot, edited_cot, examples)
                response = client.models.generate_content(
                    model="gemini-3-flash-preview",
                    config=types.GenerateContentConfig(
                        system_instruction=perturbation_instruction),
                    contents=prompt_for_input
                )

                rating = extract_rating(response.text)

                ws.cell(row=idx + 2, column=gemini_col_idx, value=rating)

        wb.save(input_file)
        print(f"  Saved to {input_file.name}")


import matplotlib.pyplot as plt

if __name__ == "__main__":
    input_folder = Path(__file__).parent / 'annotation_data'

    # Collect all rows across files and sheets
    all_rows = []

    for input_file in input_folder.glob('sqa_input_annotation.xlsx'):
        print(f"Reading file: {input_file.name}")
        sheets = pd.read_excel(input_file, sheet_name=None)

        for sheet_name, df in sheets.items():
            if 'gemini' not in df.columns:
                print(f"  'gemini' column not found in sheet {sheet_name}, skipping.")
                continue
            df['score'] = pd.to_numeric(df['score'], errors='coerce')
            df['gemini'] = pd.to_numeric(df['gemini'], errors='coerce')

            # Only keep rows where score is between 0 and 3 and gemini score exists
            mask = df['score'].between(0, 3) & df['gemini'].notna()
            all_rows.append(df[mask][['score', 'gemini']])

    # Combine all rows into one dataframe
    combined_df = pd.concat(all_rows, ignore_index=True)

    # Compute average gemini score for each true score
    avg_gemini = combined_df.groupby('score')['gemini'].mean()

    # Plot
    fig, ax = plt.subplots(figsize=(8, 6))
    bars = ax.bar(avg_gemini.index, avg_gemini.values, color=['#4C72B0', '#DD8452', '#55A868', '#C44E52'], width=0.5)

    # Add value labels on top of each bar
    for bar in bars:
        height = bar.get_height()
        ax.text(bar.get_x() + bar.get_width() / 2, height + 0.01, f'{height:.2f}', ha='center', va='bottom', fontsize=11)

    ax.set_xticks([0, 1, 2, 3])
    ax.set_xticklabels(['Score 0', 'Score 1', 'Score 2', 'Score 3'])
    ax.set_xlabel('True Score')
    ax.set_ylabel('Average Gemini Predicted Score')
    ax.set_title('Average Gemini Predicted Score per True Score Group')
    ax.set_ylim(0, max(avg_gemini.values) + 0.2)

    plt.tight_layout()
    plt.savefig(input_folder / 'gemini_scores_bar_chart.png', dpi=150, bbox_inches='tight')
    plt.show()
    print("Bar chart saved to gemini_scores_bar_chart.png")