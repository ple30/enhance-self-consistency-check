import torch
import torch.nn.functional as F
from sentence_transformers import SentenceTransformer
import glob
import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

model = SentenceTransformer("all-MiniLM-L6-v2")

def compute_cosine_distance(texts1, texts2, batch_size=8):
    if not texts1 or not texts2:
        return []

    all_distances = []
    
    # Process in smaller batches
    for i in range(0, len(texts1), batch_size):
        batch1 = texts1[i : i + batch_size]
        batch2 = texts2[i : i + batch_size]
        
        # 1. Generate embeddings (Ensure they are 2D: [batch, dims])
        # Using a typical SentenceTransformer 'model.encode' as an example
        emb1 = model.encode(batch1, convert_to_tensor=True)
        emb2 = model.encode(batch2, convert_to_tensor=True)
        
        # 2. Compute Similarity
        # dim=1 works here because emb1/emb2 are 2D matrices [N, D]
        cos_sim = F.cosine_similarity(emb1, emb2, dim=1)
        
        # 3. Convert to Distance (1 - sim)
        dist = 1 - cos_sim
        all_distances.extend(dist.cpu().tolist())
        
    return all_distances

'''def compute_cosine_distance(text1, text2):
    e1 = model.encode(text1, convert_to_tensor=True)
    e2 = model.encode(text2, convert_to_tensor=True)
    cos_sim = F.cosine_similarity(e1, e2, dim=0)
    cos_dist = ((1 - cos_sim) / 2).tolist()  
    return cos_dist'''

def add_cosine_distances_to_excel(input_file, original_col, edited_col):
    print(f"\nProcessing file: {input_file.name}")

    sheets = pd.read_excel(input_file, sheet_name=None)
    wb = load_workbook(input_file)

    for sheet_name, df in sheets.items():
        print(f"  Sheet: {sheet_name}")
        ws = wb[sheet_name]

        # Find or create the 'cosine distance' column
        header_row = [cell.value for cell in ws[1]]
        if 'cosine distance' in header_row:
            pred_col_idx = header_row.index('cosine distance') + 1
        else:
            pred_col_idx = len(header_row) + 1
            ws.cell(row=1, column=pred_col_idx, value='cosine distance')

        # Filter rows where 'score' is between 0 and 3
        mask = pd.to_numeric(df['score'], errors='coerce').between(0, 3)

        '''for idx, row in df[mask].iterrows():
            if input:
                text1 = str(row['Question'])
                text2 = str(row['Edited Question'])
            elif cot:
                text1 = str(row['Full original CoT'])
                text2 = str(row['Full edited CoT'])

            dist = compute_cosine_distance(text1, text2)
            ws.cell(row=idx + 2, column=pred_col_idx, value=dist)
        '''
        # Batch processing for efficiency
        filtered_df = df[mask]
        if filtered_df.empty:
            print(f"  No valid rows found, skipping.")
            continue

        texts1 = filtered_df[original_col].astype(str).tolist()
        texts2 = filtered_df[edited_col].astype(str).tolist()

        # 1. Get the distances
        distances = compute_cosine_distance(texts1, texts2)

        # 2. Map them back to the filtered rows specifically
        filtered_df['dist'] = distances

        # 3. Write to Excel using the index-preserved values
        for idx, row in filtered_df.iterrows():
            # idx + 2 assumes the Excel header is row 1 and DF index starts at 0
            ws.cell(row=idx + 2, column=pred_col_idx, value=row['dist'])

    wb.save(input_file)
    print(f"  Saved to {input_file.name}")
if __name__ == "__main__":
    annotation_folder = Path(__file__).parent / 'annotation_data'
    for file in annotation_folder.glob('*input_annotation.xlsx'):
        add_cosine_distances_to_excel(file, original_col='Question', edited_col='Edited Question')
    for file in annotation_folder.glob('*cot_annotation.xlsx'):
        add_cosine_distances_to_excel(file, original_col='Full original CoT', edited_col='Full edited CoT')



'''import matplotlib.pyplot as plt

if __name__ == "__main__":
    input_folder = Path(__file__).parent / 'cot_annotation_1'

    low_scores = []   # cosine distance where score = 0 or 1
    high_scores = []  # cosine distance where score = 2 or 3

    for input_file in input_folder.glob('sqa*.xlsx'):
        print(f"\nProcessing file: {input_file.name}")

        sheets = pd.read_excel(input_file, sheet_name=None)

        for sheet_name, df in sheets.items():
            print(f"  Sheet: {sheet_name}")

            for idx, row in df.iterrows():
                if row['score'] in [0, 1]:
                    low_scores.append(row['cosine distance'])
                elif row['score'] in [2, 3]:
                    high_scores.append(row['cosine distance'])

    # Draw boxplot
    plt.figure(figsize=(8, 6))
    plt.boxplot([low_scores, high_scores], labels=['Score 0-1', 'Score 2-3'])
    plt.title('Predicted Score Distribution by Score Group')
    plt.ylabel('Cosine Distance (Predicted Score)')
    plt.xlabel('Score Group')
    plt.savefig(input_folder / 'cosine_boxplot.png', dpi=150, bbox_inches='tight')
    plt.show()
    print("\nBoxplot saved to cosine_boxplot.png")
'''